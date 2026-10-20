import pandas as pd
import pytest
from openpyxl import load_workbook

from sales_report.cleaner import clean, load_all, parse_dates, parse_money, standardise_columns
from sales_report.report import build_report


def raw(rows, cols=("date", "product", "quantity", "unit_price")):
    return pd.DataFrame(rows, columns=list(cols), dtype=str)


def test_parse_dates_handles_iso_slash_and_australian_formats():
    s = pd.Series(["2026-09-01", "2026/09/02", "03/09/2026", "not a date"])
    out = parse_dates(s)
    assert list(out[:3].dt.strftime("%Y-%m-%d")) == ["2026-09-01", "2026-09-02", "2026-09-03"]
    assert pd.isna(out[3])


def test_parse_money_strips_symbols_and_commas():
    assert list(parse_money(pd.Series(["$12.50", "1,234.00", " 7 "]))) == [12.5, 1234.0, 7.0]


def test_column_aliases_are_matched_case_insensitively():
    df = pd.DataFrame(columns=["Transaction Date", "SKU Name", "Units", "Unit_Price", "Order No"])
    assert list(standardise_columns(df, "c.csv").columns) == [
        "date", "product", "quantity", "unit_price", "order_id"]


def test_missing_required_column_gives_clear_error():
    with pytest.raises(ValueError, match="unit_price"):
        standardise_columns(pd.DataFrame(columns=["date", "product", "qty"]), "bad.csv")


def test_product_names_are_normalised():
    df, _ = clean(raw([["2026-09-01", "  matcha LATTE ", "1", "7"]]), "Test")
    assert df.loc[0, "product"] == "Matcha Latte"


def test_identical_sales_without_order_id_are_both_kept_after_normalising():
    # Different spelling of the same sale = two genuine orders, not a duplicate export.
    df, stats = clean(raw([["2026-09-01", "Taro Milk Tea", "1", "7"],
                           ["2026-09-01", "taro milk tea", "1", "7"]]), "Test")
    assert len(df) == 2 and stats["dropped_duplicates"] == 0


def test_duplicate_order_ids_are_removed():
    cols = ("order_id", "date", "product", "quantity", "unit_price")
    df, stats = clean(raw([["A1", "2026-09-01", "Tea", "1", "7"],
                           ["A1", "2026-09-01", "Tea", "1", "7"],
                           ["A2", "2026-09-01", "Tea", "1", "7"]], cols), "Test")
    assert len(df) == 2 and stats["dropped_duplicates"] == 1


def test_incomplete_rows_dropped_and_refunds_reduce_revenue():
    df, stats = clean(raw([["2026-09-01", "Tea", "2", "$5"],
                           ["2026-09-01", "Tea", "", "$5"],
                           ["2026-09-02", "Tea", "-1", "$5"]]), "Test")
    assert stats["dropped_incomplete"] == 1
    assert stats["refund_rows"] == 1
    assert df["revenue"].sum() == 5.0


def test_end_to_end_report(tmp_path):
    (tmp_path / "in").mkdir()
    raw([["2026-08-01", "Tea", "2", "5"], ["2026-09-01", "Coffee", "1", "4"]]).to_csv(
        tmp_path / "in" / "store_test.csv", index=False)
    df, stats = load_all(tmp_path / "in")
    out = build_report(df, stats, tmp_path / "report.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "By Product", "Clean Data", "Data Quality Log"]
    assert wb["Summary"]["A5"].value == 14.0  # total revenue KPI
