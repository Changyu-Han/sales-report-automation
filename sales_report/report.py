"""Build a formatted Excel report with native (editable) Excel charts."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
KPI_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
MONEY = '"$"#,##0.00'


def build_summaries(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    monthly = (
        df.assign(month=df["date"].dt.to_period("M").astype(str))
        .pivot_table(index="month", columns="store", values="revenue", aggfunc="sum", fill_value=0)
        .round(2)
        .reset_index()
    )
    monthly.columns.name = None
    by_product = (
        df.groupby("product", as_index=False)
        .agg(units_sold=("quantity", "sum"), revenue=("revenue", "sum"))
        .sort_values("revenue", ascending=False)
        .round(2)
    )
    by_store = (
        df.groupby("store", as_index=False)
        .agg(orders=("revenue", "size"), units_sold=("quantity", "sum"), revenue=("revenue", "sum"))
        .sort_values("revenue", ascending=False)
        .round(2)
    )
    return {"monthly": monthly, "by_product": by_product, "by_store": by_store}


def write_table(ws: Worksheet, df: pd.DataFrame, start_row: int, start_col: int = 1,
                money_cols: tuple[str, ...] = ()) -> None:
    """Write a DataFrame as a styled table."""
    for j, col in enumerate(df.columns):
        cell = ws.cell(row=start_row, column=start_col + j,
                       value=str(col).replace("_", " ").title())
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        numeric = pd.api.types.is_numeric_dtype(df[col])
        cell.alignment = Alignment(horizontal="right" if numeric else "left")
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            cell = ws.cell(row=start_row + i, column=start_col + j, value=value)
            cell.border = Border(bottom=THIN)
            col_name = df.columns[j]
            if col_name in money_cols or (money_cols == ("*",) and j > 0):
                cell.number_format = MONEY
            elif hasattr(value, "strftime"):
                cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(horizontal="left")
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = "#,##0" if float(value).is_integer() else "#,##0.00"


def autofit(ws: Worksheet, min_width: int = 10, max_width: int = 40, skip_rows: int = 0) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=skip_rows + 1):
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(width + 2, max_width))


def summary_sheet(wb: Workbook, df: pd.DataFrame, s: dict[str, pd.DataFrame]) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Sales Report"
    ws["A1"].font = Font(size=18, bold=True, color="1F4E78")
    ws["A2"] = f"{df['date'].min():%d %b %Y} – {df['date'].max():%d %b %Y}  ·  {df['store'].nunique()} stores"
    ws["A2"].font = Font(italic=True, color="595959")

    best = s["by_product"].iloc[0]
    kpis = [
        ("Total revenue", df["revenue"].sum(), MONEY),
        ("Orders", len(df), "#,##0"),
        ("Units sold", df["quantity"].sum(), "#,##0"),
        ("Avg order value", df["revenue"].mean(), MONEY),
        ("Best seller", best["product"], None),
    ]
    for j, (label, value, fmt) in enumerate(kpis):
        col = 1 + j
        label_cell = ws.cell(row=4, column=col, value=label)
        label_cell.font = Font(color="595959")
        cell = ws.cell(row=5, column=col, value=value)
        cell.font = Font(size=14, bold=True)
        if fmt:
            cell.number_format = fmt
        for c in (label_cell, cell):  # label sits directly above its value
            c.fill = KPI_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 24

    ws["A8"] = "Revenue by store"
    ws["A8"].font = Font(size=12, bold=True)
    write_table(ws, s["by_store"], start_row=9, money_cols=("revenue",))
    n = len(s["by_store"])
    chart = BarChart()
    chart.title, chart.y_axis.title = "Revenue by store", "Revenue ($)"
    chart.add_data(Reference(ws, min_col=4, min_row=9, max_row=9 + n), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=10, max_row=9 + n))
    chart.legend, chart.height, chart.width = None, 7, 14
    ws.add_chart(chart, "F8")

    start = 12 + n + 12
    ws.cell(row=start - 1, column=1, value="Monthly revenue trend").font = Font(size=12, bold=True)
    write_table(ws, s["monthly"], start_row=start, money_cols=("*",))
    m, stores = len(s["monthly"]), len(s["monthly"].columns) - 1
    line = LineChart()
    line.title, line.y_axis.title = "Monthly revenue by store", "Revenue ($)"
    line.add_data(Reference(ws, min_col=2, max_col=1 + stores, min_row=start, max_row=start + m),
                  titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=start + 1, max_row=start + m))
    line.height, line.width = 7, 14
    for series in line.series:
        series.smooth = False
    ws.add_chart(line, f"F{start - 1}")
    autofit(ws, min_width=14, skip_rows=2)
    # KPI values use a larger font, so give those columns extra room.
    for j, (_, value, _) in enumerate(kpis):
        letter = get_column_letter(1 + j)
        needed = len(f"{value:,.2f}" if isinstance(value, float) else str(value)) * 1.4 + 4
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width, needed)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def product_sheet(wb: Workbook, by_product: pd.DataFrame) -> None:
    ws = wb.create_sheet("By Product")
    write_table(ws, by_product, start_row=1, money_cols=("revenue",))
    n = len(by_product)
    chart = BarChart()
    chart.type = "bar"  # horizontal bars read better for product names
    chart.title = "Revenue by product"
    chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=1 + n), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=1 + n))
    chart.legend, chart.height, chart.width = None, 8, 16
    ws.add_chart(chart, "E2")
    autofit(ws)


def data_sheet(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Clean Data")
    cols = [c for c in ["order_id", "date", "store", "product", "quantity", "unit_price", "revenue"]
            if c in df.columns]
    write_table(ws, df[cols], start_row=1, money_cols=("unit_price", "revenue"))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autofit(ws)


def quality_sheet(wb: Workbook, stats: list[dict]) -> None:
    ws = wb.create_sheet("Data Quality Log")
    log_df = pd.DataFrame(stats)[["file", "store", "rows_in", "dropped_incomplete",
                                   "dropped_duplicates", "refund_rows", "rows_out"]]
    write_table(ws, log_df, start_row=1)
    note = len(log_df) + 3
    ws.cell(row=note, column=1, value="Refund rows (negative quantity) are kept and reduce revenue.").font = \
        Font(italic=True, color="595959")
    autofit(ws)


def build_report(df: pd.DataFrame, stats: list[dict], output: Path) -> Path:
    summaries = build_summaries(df)
    wb = Workbook()
    summary_sheet(wb, df, summaries)
    product_sheet(wb, summaries["by_product"])
    data_sheet(wb, df)
    quality_sheet(wb, stats)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output
